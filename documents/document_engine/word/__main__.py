"""A check of the Word tab without a model:  cd backend && python -m document_engine.word [out.docx]"""

import re
import sys
from io import BytesIO

import docx
from openpyxl import Workbook
from ollama_pipeline import UnsupportedDocument

from .build import build_docx
from .service import share
from .sources import read_source

SAMPLE = """---
title: "Quarterly review"
subtitle: "Sales and costs"
---

# Summary

Revenue grew **24%**; margin is ==6 pp== higher.

::: callout type="warning" title="Risk"
One supplier makes up half of the costs.
:::

| Region | Revenue |
|--------|---------|
| North | 1,200 |
| South | 950 |

- [ ] Review supplier contracts

::: signature name="A. Rao" role="Finance lead" date="1 Oct 2026"
:::
"""


def main() -> None:
    book = Workbook()
    book.active.title = "Sales"
    for row in (("Region", "Revenue"), ("North", 1200.0), (None, None), ("South", 950.5)):
        book.active.append(row)
    saved = BytesIO()
    book.save(saved)
    text, truncated = read_source("sales.xlsx", saved.getvalue())
    assert "Sheet: Sales" in text and "| North | 1200 |" in text and "| South | 950.5 |" in text and not truncated, text
    assert read_source("notes.txt", b"hello")[0] == "hello"
    for bad in ("old.xls", "photo.png"):
        try:
            read_source(bad, b"x")
            raise AssertionError(f"{bad} was accepted")
        except UnsupportedDocument:
            pass

    joined, cut = share([{"name": "long", "content": "a" * 900}, {"name": "short", "content": "b" * 100}], 500)
    assert cut == ["long"] and "b" * 100 in joined and joined.count("a") == 400, (cut, joined.count("a"))

    data, name = build_docx(SAMPLE, {"template": "report", "accent": "#C0392B", "toc": True})
    document = docx.Document(BytesIO(data))
    words = re.findall(r"<w:t[^>]*>([^<]*)</w:t>", document.element.xml)  # also the cover's nested tables
    assert name == "Quarterly review", name
    for expected in ("Quarterly review", "Sales and costs", "Contents", "Summary", "North", "One supplier makes up half of the costs.", "A. Rao"):
        assert expected in words, (expected, words)
    assert not any(w.startswith("Right-click") for w in words), "the contents page still shows Word's instruction"
    assert "C0392B" in document.element.xml, "the accent colour is not in the file"

    plain, title = build_docx("# Only a heading\n\nSome text.", {"template": "memo", "cover": False})
    assert title == "Only a heading" and "Only a heading" in "\n".join(p.text for p in docx.Document(BytesIO(plain)).paragraphs)

    # A letterhead keeps its header; a file whose body is the design is refused.
    head = docx.Document()
    head.sections[0].header.paragraphs[0].text = "ACME LTD"
    saved = BytesIO()
    head.save(saved)
    on_head, _ = build_docx(SAMPLE, {"template": "letter"}, ("acme.docx", saved.getvalue()))
    result = docx.Document(BytesIO(on_head))
    assert result.sections[0].header.paragraphs[0].text == "ACME LTD" and any("Summary" in p.text for p in result.paragraphs)

    if len(sys.argv) > 1:
        with open(sys.argv[1], "wb") as out:
            out.write(data)
    print("Word tab: sources, sharing, building and letterhead all pass.")


if __name__ == "__main__":
    main()
