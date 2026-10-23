"""The bundled slides tool (tools/generate_slides.py) as Admin Panel -> Tools runs it. Run from backend/:

    .venv/bin/python tests/slides_tool.py

No server, no model: the app as main.py builds it, on a temporary auth.db and a temporary tool_files
folder. The tool is saved through the admin route and called the way the chat calls it. httpx is replaced
by a recorder, so no request leaves this process and every one the tool tries is seen.
"""
import asyncio, base64, json, re, sys, tempfile, time
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))
import httpx
from fastapi import HTTPException
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from app import model_settings, tools
from app.auth import current_user, init_db
from app.main import create_app

results = []
def check(name, ok):
    results.append(bool(ok)); print(("PASS " if ok else "FAIL ") + name)


tmp = Path(tempfile.mkdtemp())
init_db(tmp / "auth.db")
model_settings.init()
tools.init()
tools.FILES_DIR = tmp / "tool_files"

# Every httpx request the tool makes lands here; `answer` decides the reply.
sent = []
def answer(request):
    if request.url.host == "api.unsplash.com":
        return httpx.Response(200, json={"urls": {"regular": unsplash_next[0]}})
    if request.url.host.endswith(".unsplash.com"):
        return httpx.Response(200, content=png(800, 600), headers={"content-type": "image/png"})
    return httpx.Response(200, content=png(800, 600))
unsplash_next = ["http://127.0.0.1/evil"]


def recording(cls):
    class Recorder(cls):
        def __init__(self, *a, **kw):
            kw["transport"] = (httpx.MockTransport(lambda r: (sent.append(str(r.url)), answer(r))[1]))
            super().__init__(*a, **kw)
    return Recorder

httpx.AsyncClient, httpx.Client = recording(httpx.AsyncClient), recording(httpx.Client)
import urllib.request
def no_urlopen(url, *a, **k):
    sent.append(f"urlopen {url}"); raise OSError("blocked")
urllib.request.urlopen = no_urlopen


def png(w=64, h=48):
    buf = BytesIO(); Image.new("RGB", (w, h), (200, 80, 40)).save(buf, "PNG"); return buf.getvalue()


admin = {"id": "a", "email": "admin@test.com", "name": "A", "role": "admin"}
user = {**admin, "id": "u", "email": "user@test.com", "role": "user"}
who = {"user": admin}


def signed_in():
    if who["user"] is None:
        raise HTTPException(401, "Not signed in.")
    return who["user"]

app = create_app()
app.dependency_overrides[current_user] = signed_in
c = TestClient(app)
A = "/api/admin/tools"
SRC = (Path(__file__).parents[2] / "tools" / "generate_slides.py").read_text()
BODY = {"id": "generate_slides", "name": "Generate Slides", "content": SRC, "meta": {"description": "pptx"}}

# --- who may install it
for who["user"], want, label in ((user, 403, "non-admin"), (None, 401, "signed out")):
    check(f"{label} create -> {want}", c.post(f"{A}/create", json=BODY).status_code == want)
who["user"] = admin
check("non-admin/signed out stored nothing", c.get(f"{A}/generate_slides").status_code == 404)
r = c.post(f"{A}/create", json=BODY)
check(f"admin create -> 200 ({r.status_code} {r.text[:300] if r.status_code != 200 else ''})", r.status_code == 200)
if r.status_code != 200:
    print(f"passed {sum(results)} of {len(results)}"); sys.exit(1)
check(f"B1: the bundled tool ({len(SRC):,} chars) is over the old 500,000 cap, so the create check above proves the raise", len(SRC) > 500_000)
r = c.post(f"{A}/create", json={**BODY, "id": "toobig", "content": "#" * 5_000_000})
check("B1: a 5 MB body is still refused (the cap was raised, not removed)", r.status_code == 422 and c.get(f"{A}/toobig").status_code == 404)

# --- what the model sees
specs = {s["name"]: s for s in c.get(f"{A}/generate_slides").json()["specs"]}
check(f"specs: generate_slides, edit_presentation, inspect_slides ({sorted(specs)})",
      set(specs) == {"generate_slides", "edit_presentation", "inspect_slides"})
params = {n: sorted(s["parameters"].get("properties", {})) for n, s in specs.items()}
check(f"specs: only model-visible params {params}",
      params.get("generate_slides") == ["content"] and params.get("edit_presentation") == ["content"]
      and params.get("inspect_slides") == ["file_id"])
check("specs: no __dunder__ param anywhere", "__" not in json.dumps(params))
check("specs: the layout list reaches the model", "AVAILABLE LAYOUTS" in specs["generate_slides"].get("description", ""))
vspec = c.get(f"{A}/generate_slides/valves/spec").json().get("properties", {})
check("valves: pptx_export_dir gone", "pptx_export_dir" not in vspec)
check("valves: template_mode_enabled, presentation_edit_enabled, inspect_slides_enabled default false",
      all(vspec.get(k, {}).get("default") is False for k in ("template_mode_enabled", "presentation_edit_enabled", "inspect_slides_enabled")))


def valves(**kw):
    r = c.post(f"{A}/generate_slides/valves/update", json=kw)
    assert r.status_code == 200, r.text


def call(fn, **kw):
    return asyncio.run(tools.get_tools(["generate_slides"], admin)[fn]["call"](**kw))


def link(res):
    m = re.search(r"\]\((/api/tools/files/([^)]+))\)", res or "")
    return (m.group(1), m.group(2)) if m else (None, None)


def stored(n):
    """Where the tool put deck n: FILES_DIR/<user id>/<n> (B3), or loose in FILES_DIR before the tool moved."""
    return next((p for p in (tools.FILES_DIR / admin["id"] / n, tools.FILES_DIR / n) if p.is_file()), None)


def deck(url):
    d = c.get(url)
    return Presentation(BytesIO(d.content)) if d.status_code == 200 else None


B64 = "data:image/png;base64," + base64.b64encode(png()).decode()
CARD = {"heading": "Plan", "icon": "target", "description": "What we do", "points": ["a", "b"]}
STEPS = [{"when": "Q1", "title": "Start", "description": "Kick-off"}, {"when": "Q2", "title": "Build", "description": "Make it"},
         {"when": "Q3", "title": "Ship", "description": "Release"}]
ITEMS = [{"icon": i, "title": i.title(), "description": "Short text"} for i in ("rocket", "shield", "users", "cloud")]
NODES = [{"label": f"Stage {i}", "description": "d"} for i in range(1, 5)]
SLIDES = [
    {"layout": "cover", "title": "Every layout", "subtitle": "QA deck", "author": "QA", "eyebrow": "TEST",
     "icon": "rocket", "date": "Oct 2026", "chips": ["one", "two"]},
    {"layout": "section", "number": "01", "eyebrow": "PART I", "title": "Context", "lead": "Why"},
    {"layout": "title_bullets", "title": "Bullets", "eyebrow": "PART I", "bullets": ["One", "Two", "Three"]},
    {"layout": "title_body", "title": "Body", "body": "First paragraph.\nSecond paragraph."},
    {"layout": "two_column_text", "title": "Two columns", "left": CARD, "right": {**CARD, "highlight": True, "badge": "Most chosen"}},
    {"layout": "comparison_two", "title": "Compare", "columns": [CARD, CARD]},
    {"layout": "kpi_row", "title": "KPIs", "stats": [{"value": "42%", "label": "Growth", "change": "+5"}, {"value": "1.2M", "label": "Users"}]},
    {"layout": "timeline_horizontal", "title": "Timeline", "steps": STEPS},
    {"layout": "process_flow", "title": "Process", "steps": STEPS},
    {"layout": "icon_list_vertical", "title": "Icon list", "items": ITEMS[:3]},
    {"layout": "icon_grid_2x2", "title": "Grid 2x2", "items": ITEMS},
    {"layout": "icon_grid_3", "title": "Grid 3", "items": ITEMS[:3]},
    {"layout": "pillars", "title": "Pillars", "items": ITEMS[:3]},
    {"layout": "chart", "title": "Bar", "chart_type": "bar", "labels": ["A", "B", "C"], "values": [3, 5, 2],
     "insight_title": "Takeaway", "insight": ["B leads"]},
    {"layout": "chart", "title": "Line", "chart_type": "line", "labels": ["Q1", "Q2", "Q3"],
     "datasets": [{"label": "2025", "data": [1, 2, 3]}, {"label": "2026", "data": [2, 3, 5]}]},
    {"layout": "chart", "title": "Area", "chart_type": "area", "labels": ["a", "b", "c"], "values": [1, 4, 2]},
    {"layout": "chart", "title": "Pie", "chart_type": "pie", "labels": ["x", "y"], "values": [60, 40]},
    {"layout": "chart", "title": "Doughnut", "chart_type": "doughnut", "labels": ["x", "y", "z"], "values": [1, 2, 3]},
    {"layout": "chart", "title": "Radar", "chart_type": "radar", "labels": ["s", "p", "q", "r"], "values": [3, 4, 2, 5]},
    {"layout": "chart", "title": "Stacked", "chart_type": "stacked_bar", "labels": ["A", "B"],
     "datasets": [{"label": "x", "data": [1, 2]}, {"label": "y", "data": [2, 1]}]},
    {"layout": "funnel", "title": "Funnel", "nodes": NODES},
    {"layout": "pyramid", "title": "Pyramid", "nodes": NODES[:3]},
    {"layout": "cycle", "title": "Cycle", "nodes": NODES},
    {"layout": "quadrant", "title": "Quadrant", "nodes": NODES, "x_axis": "Effort", "y_axis": "Impact",
     "points": [{"label": "P1", "x": 0.2, "y": 0.8}]},
    {"layout": "bullseye", "title": "Bullseye", "nodes": NODES[:3]},
    {"layout": "quote", "quote": "Less is more.", "author": "Mies", "role": "Architect"},
    {"layout": "alert", "title": "Heads up", "level": "warning", "bullets": ["Check this", "And this"]},
    {"layout": "table", "title": "Table", "headers": ["Name", "Value"], "rows": [["a", "1"], ["b", "2"]]},
    {"layout": "text_image_right", "title": "Image right", "bullets": ["x", "y"], "base64": B64},
    {"layout": "image_left_text_right", "title": "Image left", "body": "Text", "base64": B64},
    {"layout": "image_full_caption", "title": "Full image", "subtitle": "Caption", "base64": B64},
    {"layout": "closing", "title": "Thanks", "eyebrow": "END", "takeaways": ["One", "Two"], "contact": "qa@test.com"},
]
layouts_doc = set(re.findall(r'"([a-z0-9_]+)"', specs["generate_slides"]["description"].split("AVAILABLE LAYOUTS", 1)[1].split("DESIGN GUIDELINES", 1)[0]))
layouts_doc -= {"bar", "line", "area", "pie", "doughnut", "radar", "stacked_bar", "info", "tip", "warning", "danger",
                "stock query", "01", "Most chosen"}
check(f"deck covers every layout in the docstring (missing {sorted(layouts_doc - {s['layout'] for s in SLIDES})})",
      layouts_doc <= {s["layout"] for s in SLIDES})

# --- every layout builds, downloads, opens
sent.clear()
res = call("generate_slides", content=json.dumps({"title": "QA Deck: every layout!", "theme": "auto", "footer": "QA", "slides": SLIDES}))
url, name = link(res)
check(f"generate: answers a /api/tools/files link ({res[:300] if not url else name})", url)
check("generate: stored name is <slug>_<32 hex>.pptx", name and re.fullmatch(r"[a-z0-9-]+_[0-9a-f]{32}\.pptx", name))
check("generate: file is in FILES_DIR/<user id>/ (B3: only its owner can download it)", name and (tools.FILES_DIR / admin["id"] / name).is_file())
prs = deck(url) if url else None
check("generate: downloads (200) and python-pptx opens it", prs is not None)
shapes = [sh for s in prs.slides for sh in s.shapes] if prs else []
check(f"generate: slide count = {len(SLIDES)} ({len(prs.slides) if prs else '-'})", prs and len(prs.slides) == len(SLIDES))
check(f"generate: native charts present ({sum(sh.has_chart for sh in shapes)})", sum(sh.has_chart for sh in shapes) >= 7)
check("generate: a native table present", any(sh.has_table for sh in shapes))
check("generate: base64 pictures placed", sum(sh.shape_type == 13 for sh in shapes) >= 3)
check("generate: no network used", sent == [])

# --- B2: a big deck must not freeze the event loop (the chat, /api/health, other chats keep running)
def loop_stall(make):
    """Longest gap between 50 ms ticks of a heartbeat task while make() runs; returns (result, gap)."""
    async def main():
        gaps, last, stop = [], [0.0], [False]
        async def beat():
            last[0] = time.perf_counter()
            while not stop[0]:
                await asyncio.sleep(0.05); now = time.perf_counter(); gaps.append(now - last[0] - 0.05); last[0] = now
        t = asyncio.create_task(beat()); await asyncio.sleep(0.2)
        out = await make(); stop[0] = True; await t
        return out, max(gaps)
    return asyncio.run(main())

big = {"title": "Stall", "slides": [{"layout": "chart", "title": f"C{i}", "chart_type": "bar", "labels": list("abcdef"),
                                     "values": [1, 2, 3, 4, 5, 6]} for i in range(250)]}
_, gap = loop_stall(lambda: tools.get_tools(["generate_slides"], admin)["generate_slides"]["call"](content=json.dumps(big)))
check(f"B2: a 250-slide deck does not stall the event loop (longest stall {gap:.2f} s, limit 0.5)", gap < 0.5)

# --- bad input -> an error text, not a crash
for bad in ("{nope", "[]", "", '"just a string"'):
    try:
        out = call("generate_slides", content=bad)
        ok = isinstance(out, str) and "/api/tools/files/" not in out and "couldn't" in out.lower()
    except Exception as exc:
        out, ok = repr(exc), False
    check(f"generate {bad!r}: error text, no crash", ok)

# --- no SSRF: a URL from the model is never fetched
sent.clear()
evil = {"title": "SSRF", "slides": [
    {"layout": "text_image_right", "title": "a", "bullets": ["x"], "image_url": "http://127.0.0.1:9/x"},
    {"layout": "image_left_text_right", "title": "b", "body": "x", "url": "http://169.254.169.254/latest/meta-data/"},
    {"layout": "image_full_caption", "title": "c", "src": "http://localhost:8000/api/auth/me"},
    {"layout": "image_full_caption", "title": "d", "image": "http://10.0.0.1/x.png", "images": ["http://10.0.0.2/y.png"]},
    {"layout": "image_grid", "title": "e", "images": ["http://127.0.0.1:5678/", {"href": "file:///etc/passwd"}]},
    {"layout": "text_image_right", "title": "f", "bullets": ["x"], "image_hint": "mountains"}]}
res = call("generate_slides", content=json.dumps(evil))
check(f"image URLs from the model: no request made ({sent})", sent == [])
check("image URLs from the model: deck still built", link(res)[0])

# Unsplash only with a key, only api.unsplash.com, and its photo only from an unsplash.com host.
valves(unsplash_access_key="KEY-123")
sent.clear(); unsplash_next[0] = "http://127.0.0.1/evil"
res = call("generate_slides", content=json.dumps({"title": "U", "slides": [{"layout": "text_image_right", "title": "a", "bullets": ["x"], "image_hint": "sea"}]}))
check(f"unsplash: api asked, its non-unsplash photo URL not fetched ({sent})",
      len(sent) == 1 and sent[0].startswith("https://api.unsplash.com/") and link(res)[0])
check("unsplash: key never in the reply", "KEY-123" not in res)
for host in ("http://images.unsplash.com.evil.test/x", "https://evil.test/?.unsplash.com", "https://unsplash.com@127.0.0.1/x"):
    sent.clear(); unsplash_next[0] = host
    valves(unsplash_access_key="KEY-123")
    call("generate_slides", content=json.dumps({"title": "U", "slides": [{"layout": "text_image_right", "title": "a", "image_hint": "sea"}]}))
    check(f"unsplash: lookalike host {host} not fetched ({sent})", len(sent) == 1 and sent[0].startswith("https://api.unsplash.com/"))
valves(unsplash_access_key="KEY-123")
sent.clear(); unsplash_next[0] = "https://images.unsplash.com/photo-1?w=1080"
res = call("generate_slides", content=json.dumps({"title": "U", "slides": [{"layout": "text_image_right", "title": "a", "bullets": ["x"], "image_hint": "sea"}]}))
p = deck(link(res)[0]) if link(res)[0] else None
check(f"unsplash: a real images.unsplash.com photo is fetched and placed ({sent})",
      len(sent) == 2 and sent[1].startswith("https://images.unsplash.com/") and p and any(sh.shape_type == 13 for sh in p.slides[0].shapes))
valves()

# --- edit_presentation reads only stored decks
edit = lambda fid, ops=(): call("edit_presentation", content=json.dumps({"presentation_file_id": fid, "operations": list(ops) or [{"op": "set_text", "slide": 1, "shape_id": 2, "text": "x"}]}))
out = edit(name)
check("edit: off by default, says so", "presentation_edit_enabled" in out and "/api/tools/files/" not in out)
valves(presentation_edit_enabled=True)
secret = tools.FILES_DIR.parent / "secret.pptx"  # outside FILES_DIR: no edit or inspect may read or change it
secret.write_bytes(stored(name).read_bytes()); secret_before = secret.read_bytes()
for bad in ("../auth.db", "../secret.pptx", "..%2Fauth.db", "/etc/passwd", str(tmp / "auth.db"), "auth.db", "a/../../auth.db",
            "nothing_0123456789abcdef0123456789abcdef.pptx", "", "x" * 300):
    out = edit(bad)
    check(f"edit: {bad[:40]!r} refused", isinstance(out, str) and "/api/tools/files/" not in out)
check("edit: the file outside FILES_DIR is untouched", secret.read_bytes() == secret_before)
first = next(sh for sh in Presentation(str(stored(name))).slides[0].shapes if sh.has_text_frame and sh.text_frame.text.strip())
out = edit(name, [{"op": "set_text", "slide": 1, "shape_id": first.shape_id, "text": "Edited title"}])
eurl, ename = link(out)
ep = deck(eurl) if eurl else None
check(f"edit: stored deck edited, new link ({out[:200] if not eurl else ename})", ep and ename != name)
check("edit: the text changed, slide count kept", ep and any(sh.has_text_frame and sh.text_frame.text == "Edited title" for sh in ep.slides[0].shapes)
      and len(ep.slides) == len(SLIDES))
valves()

# --- template mode / inspect: not available here
for on in (False, True):
    valves(template_mode_enabled=on, inspect_slides_enabled=on)
    if not on:  # the default; with the valves on, template mode works on decks this tool stored
        out = call("inspect_slides", file_id=name)
        check("inspect_slides (valves off): 'not available' error", "not available" in out.lower() and "/api/tools/files/" not in out)
        out = call("generate_slides", content=json.dumps({"title": "T", "reference_file_id": name, "slides": [{"layout": "cover", "title": "x"}]}))
        check("reference_file_id (valves off): no template read", "not available" in out.lower() or "ignored" in out.lower())
    for bad in ("../secret.pptx", "../auth.db", "/etc/passwd"):
        out = call("inspect_slides", file_id=bad)
        check(f"inspect_slides {bad!r} (valves {on}): refused, nothing inventoried", '"ok": true' not in out.lower() and "/api/tools/files/" not in out)
    out = call("generate_slides", content=json.dumps({"title": "T", "reference_file_id": "../auth.db", "slides": [{"reuse": {"slide": 0}}]}))
    check(f"reuse with ../auth.db (valves {on}): error, no file", "/api/tools/files/" not in out)
valves()
check("template/inspect: the file outside FILES_DIR is untouched", secret.read_bytes() == secret_before)

# --- the download route
check("download: traversal -> 404", c.get("/api/tools/files/..%2Fauth.db").status_code == 404)
who["user"] = user
check("download: another signed-in user with the link -> 404 (B3)", c.get(url).status_code == 404)
who["user"] = None
check("download: signed out -> 401", c.get(url).status_code == 401)
who["user"] = admin
check("no file written outside FILES_DIR", {p.name for p in tmp.iterdir()} <= {"auth.db", "auth.db-wal", "auth.db-shm", "tool_files", "secret.pptx"})

print(f"passed {sum(results)} of {len(results)}")
sys.exit(0 if all(results) else 1)
